"""投标文件格式章节化 docx（Task 12）。

内容来源统一为 MinerU content_list（方案 A'，经用户 2026-09-29 确认）：
章节边界 100% 复用 Task 8 chapters.json 的页码范围，三种源文件
（文字版 PDF / 扫描件 PDF / docx）走同一切分路径，不直读原文件。

分层：章节定位/切分计划为纯函数（离线可单测），python-docx 落盘为唯一副作用。
产物：``<project>/parse/docx/<source_stem>/NN_章节名.docx`` + ``manifest.json``。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import UTC, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn

from app.parse import paths

# 正文中文字体（与项目文档字体偏好一致：仿宋）
_BODY_FONT = "仿宋"
# 章标题归一化匹配关键词
_FORMAT_KEYWORD = "投标文件格式"
_COVER_TITLE = "封面"
# EXTERNAL 标记关键词：系统生成/报价类章节，展示但标注"系统外/不制作"，不参与后续素材/模板/渲染
_EXTERNAL_KEYWORDS = (
    "开标一览表",
    "投标报价",
    "分项报价表",
    "报价表",
    "报价函",
    "投标函附录",
    "法定代表人授权书",  # 部分招标文件视作系统模板，按外部处理
)


@dataclass
class SectionPlan:
    """单个逐章 docx 的切分计划（页码闭区间）。"""

    source_stem: str
    seq: int  # 文件序号：封面恒 0，正文从 1 起
    title: str
    kind: str  # "cover" | "section"
    start_page: int
    end_page: int

    @property
    def checkpoint_key(self) -> str:
        return f"docx:{self.source_stem}:{self.seq:02d}"

    @property
    def filename(self) -> str:
        return f"{self.seq:02d}_{paths.safe_filename(self.title)}.docx"

    @property
    def is_external(self) -> bool:
        """是否为系统生成/报价类章节（不参与后续素材/模板/渲染环节）。"""
        return any(kw in self.title for kw in _EXTERNAL_KEYWORDS)


@dataclass
class SourcePlan:
    """单个源文件的格式章节切分计划。"""

    source_stem: str
    found: bool
    sections: list[SectionPlan]
    note: str | None = None


# ---------- 章节定位与切分计划（纯逻辑） ----------


def _norm(text: str) -> str:
    return re.sub(r"[\s　.．。:：、,，;；()（）]", "", text)


def _iter_nodes(nodes: list[dict[str, Any]]):
    """子树前序遍历。"""
    for node in nodes:
        yield node
        children = node.get("children") or []
        yield from _iter_nodes(children)


def find_format_chapter(chapters: list[dict[str, Any]]) -> dict[str, Any] | None:
    """在章节树中定位"投标文件格式"章（标题归一化包含关键词）。"""
    for node in _iter_nodes(chapters):
        if _FORMAT_KEYWORD in _norm(str(node.get("title", ""))):
            return node
    return None


def find_cover_in_subtree(node: dict[str, Any]) -> dict[str, Any] | None:
    """在格式章子树内定位"封面"节点（精确标题匹配）。"""
    for sub in _iter_nodes([node]):
        if _norm(str(sub.get("title", ""))) == _COVER_TITLE:
            return sub
    return None


def _end_page(node: dict[str, Any], fallback: int) -> int:
    end = node.get("end_page")
    if end is None:
        return fallback
    return max(int(node["start_page"]), int(end))


def _fallback_headings(
    content_list: list[dict[str, Any]], start_page: int, end_page: int
) -> list[dict[str, Any]]:
    """格式章无 children 时：按 MinerU text_level（≤2）在页码范围内补识别子节标题。"""
    headings: list[dict[str, Any]] = []
    for block in content_list:
        page = int(block.get("page_idx", 0))
        if not (start_page <= page <= end_page):
            continue
        level = block.get("text_level")
        text = str(block.get("text", "")).strip()
        if not text or not isinstance(level, int) or level > 2:
            continue
        if headings and headings[-1]["page_idx"] == page and headings[-1]["text"] == text:
            continue
        headings.append({"text": text, "page_idx": page})
    return headings


def plan_source(source: dict[str, Any], content_list: list[dict[str, Any]] | None) -> SourcePlan:
    """为单个源文件制定逐章 docx 计划。

    source：chapters.json 的单个 sources 元素（含 chapters/page_count）。
    content_list：该源文件 MinerU content_list 块数组。
    """
    stem = str(source.get("source_stem", ""))
    chapters = source.get("chapters") or []
    page_count = max(int(source.get("page_count", 1)) - 1, 0)
    fmt = find_format_chapter(chapters)
    if fmt is None:
        return SourcePlan(stem, found=False, sections=[], note="未识别到「投标文件格式」章节")

    blocks = content_list or []
    fmt_start = int(fmt["start_page"])
    fmt_end = _end_page(fmt, page_count)
    sections: list[SectionPlan] = []

    cover_node = find_cover_in_subtree(fmt)
    # chapters.py 将"封面"识别为一级章（常为格式章的同级而非子节点），子树找不到时搜整棵树
    if cover_node is None:
        for top in chapters:
            cover_node = find_cover_in_subtree(top)
            if cover_node is not None:
                break
    cover_is_direct_child = False
    child_nodes = list(fmt.get("children") or [])
    if cover_node is not None:
        cover_is_direct_child = any(c is cover_node for c in child_nodes)
        sections.append(
            SectionPlan(
                source_stem=stem,
                seq=0,
                title=_COVER_TITLE,
                kind="cover",
                start_page=int(cover_node["start_page"]),
                end_page=_end_page(cover_node, fmt_end),
            )
        )

    body_children = [c for c in child_nodes if not (cover_is_direct_child and c is cover_node)]
    if body_children:
        # 编号正文：直接子章各成一个 docx（嵌套子节以标题形式保留在同一 docx 内）
        seq = 1
        for child in body_children:
            sections.append(
                SectionPlan(
                    source_stem=stem,
                    seq=seq,
                    title=str(child.get("title", f"未命名章节{seq}")),
                    kind="section",
                    start_page=int(child["start_page"]),
                    end_page=_end_page(child, fmt_end),
                )
            )
            seq += 1
    else:
        headings = _fallback_headings(blocks, fmt_start, fmt_end)
        # 排除格式章自身的标题块（与章同页同名，不是子章节）
        fmt_title = _norm(str(fmt.get("title", "")))
        headings = [
            h
            for h in headings
            if not (int(h["page_idx"]) == fmt_start and _norm(str(h["text"])) == fmt_title)
        ]
        if headings:
            seq = 1
            for i, head in enumerate(headings):
                h_start = int(head["page_idx"])
                h_end = int(headings[i + 1]["page_idx"]) - 1 if i + 1 < len(headings) else fmt_end
                sections.append(
                    SectionPlan(
                        source_stem=stem,
                        seq=seq,
                        title=str(head["text"]),
                        kind="section",
                        start_page=max(fmt_start, min(h_start, fmt_end)),
                        end_page=max(h_start, min(h_end, fmt_end)),
                    )
                )
                seq += 1
        else:
            # 章内无任何子节结构：整章合并为一个 docx
            sections.append(
                SectionPlan(
                    source_stem=stem,
                    seq=1,
                    title=str(fmt.get("title", _FORMAT_KEYWORD)),
                    kind="section",
                    start_page=fmt_start,
                    end_page=fmt_end,
                )
            )

    return SourcePlan(stem, found=True, sections=sections)


def plan_all_sources(
    chapters_payload: dict[str, Any],
    content_lists: dict[str, list[dict[str, Any]]],
) -> list[SourcePlan]:
    plans: list[SourcePlan] = []
    for source in chapters_payload.get("sources", []):
        stem = str(source.get("source_stem", ""))
        plans.append(plan_source(source, content_lists.get(stem)))
    return plans


def collect_blocks(
    content_list: list[dict[str, Any]], section: SectionPlan
) -> list[dict[str, Any]]:
    """收集章节页码闭区间内的块；跳过与章节标题同页同名的标题块（避免标题重复）。"""
    target_title = _norm(section.title)
    blocks: list[dict[str, Any]] = []
    for block in content_list:
        page = int(block.get("page_idx", 0))
        if not (section.start_page <= page <= section.end_page):
            continue
        text = str(block.get("text", "")).strip()
        if (
            page == section.start_page
            and isinstance(block.get("text_level"), int)
            and _norm(text) == target_title
        ):
            continue
        blocks.append(block)
    return blocks


# ---------- python-docx 落盘（副作用层） ----------


class _TableHTMLParser(HTMLParser):
    """MinerU table_body（HTML <table>）→ 行/列文本。"""

    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[str]] = []
        self._in_cell = False
        self._cell_parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "tr":
            self.rows.append([])
        elif tag in ("td", "th"):
            self._in_cell = True
            self._cell_parts = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._in_cell:
            cell = "".join(self._cell_parts).strip()
            if self.rows:
                self.rows[-1].append(cell)
            self._in_cell = False

    def handle_data(self, data: str) -> None:
        if self._in_cell:
            self._cell_parts.append(data)


def _parse_html_table(html: str) -> list[list[str]]:
    parser = _TableHTMLParser()
    parser.feed(html)
    return [row for row in parser.rows if any(cell for cell in row)]


def _parse_md_table(text: str) -> list[list[str]]:
    """老版本 MinerU Markdown 表格降级解析（|-|-| 分隔行丢弃）。"""
    rows: list[list[str]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def _set_body_font(document: DocumentObject) -> None:
    style = document.styles["Normal"]
    style.font.name = _BODY_FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), _BODY_FONT)


def _add_table(document: DocumentObject, rows: list[list[str]]) -> None:
    col_count = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(col_count):
            table.cell(i, j).text = row[j] if j < len(row) else ""


def blocks_to_docx(blocks: list[dict[str, Any]], out_path: Path, doc_title: str) -> dict[str, int]:
    """块数组 → docx 文件。返回 {paragraphs,tables,images,red_flags} 计数。"""
    document = Document()
    _set_body_font(document)
    document.add_heading(doc_title, level=0)

    counters = {"paragraphs": 0, "tables": 0, "images": 0, "red_flags": 0}

    for block in blocks:
        btype = block.get("type")
        if btype == "text":
            text = str(block.get("text", "")).strip()
            if not text:
                continue
            level = block.get("text_level")
            if isinstance(level, int):
                document.add_heading(text, level=min(level, 9))
            else:
                document.add_paragraph(text)
                counters["paragraphs"] += 1
        elif btype == "table":
            rows: list[list[str]] = []
            body_html = block.get("table_body")
            if isinstance(body_html, str) and "<" in body_html:
                rows = _parse_html_table(body_html)
            if not rows:
                rows = _parse_md_table(str(block.get("text", "")))
            if rows:
                _add_table(document, rows)
                counters["tables"] += 1
            else:
                raw_excerpt = str(block.get("text", ""))[:80]
                document.add_paragraph(f"【表格无法还原，请对照原文件：{raw_excerpt}】")
                counters["red_flags"] += 1
        elif btype == "image":
            caption = block.get("img_caption") or block.get("image_caption") or "图片内容"
            document.add_paragraph(f"【{caption}，请对照原文件】")
            counters["images"] += 1
            counters["red_flags"] += 1

    if not blocks:
        document.add_paragraph("【该章节页码范围内未识别到内容，请对照原文件核对】")
        counters["red_flags"] += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    return counters


def render_section(
    section: SectionPlan,
    content_lists: dict[str, list[dict[str, Any]]],
    project_dir: Path,
    docx_root: Path,
) -> dict[str, Any]:
    """渲染单个章节 docx，返回供 checkpoint output/manifest 使用的记录。"""
    blocks = collect_blocks(content_lists.get(section.source_stem, []), section)
    source_dir = paths.ensure_within_project(
        project_dir, docx_root / paths.safe_filename(section.source_stem)
    )
    out_path = source_dir / section.filename
    counters = blocks_to_docx(blocks, out_path, section.title)
    return {
        "source_stem": section.source_stem,
        "seq": section.seq,
        "title": section.title,
        "kind": section.kind,
        "file": str(out_path.relative_to(docx_root)).replace("\\", "/"),
        "start_page": section.start_page,
        "end_page": section.end_page,
        **counters,
    }


def assemble_manifest(source_plans: list[SourcePlan], generated_at: str) -> dict[str, Any]:
    """组装 manifest 骨架（文件状态由编排层按 checkpoint 结果回填）。"""
    return {
        "generated_at": generated_at,
        "completed": False,
        "total_files": sum(len(p.sections) for p in source_plans),
        "sources": [
            {
                "source_stem": p.source_stem,
                "found": p.found,
                "cover": any(s.kind == "cover" for s in p.sections),
                "note": p.note,
                "sections": [
                    {
                        "seq": s.seq,
                        "title": s.title,
                        "kind": s.kind,
                        "checkpoint_key": s.checkpoint_key,
                        "is_external": s.is_external,
                    }
                    for s in p.sections
                ],
            }
            for p in source_plans
        ],
    }


def now_iso() -> str:
    return datetime.now(UTC).isoformat()
