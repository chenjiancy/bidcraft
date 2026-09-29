"""Task 12 投标文件格式章节化 docx 集成测试（TR-12.1～12.7）。

MinerU 以 fake 替换，content_list 构造含"封面"+"投标文件格式"章及其子节，
验证三来源（文字版 PDF/扫描件 PDF/docx）统一路径、命名、隔离、按章失败恢复。
"""

import json
from pathlib import Path
from typing import Any

import pytest
from fastapi.testclient import TestClient

from app.parse import docx_build
from app.parse import service as parse_service
from app.parse.engine import RawResult
from tests.integration.test_parse_api import FakeMineru, _patch_engine, make_text_pdf


def _format_blocks(_stem: str) -> list[dict[str, Any]]:
    """含封面 + 投标文件格式章（2 个子节 + 1 个表格块）的 content_list。"""
    return [
        {"type": "text", "text": "监理招标文件", "page_idx": 0, "text_level": 1},
        {"type": "text", "text": "投标人须知前附表内容", "page_idx": 1},
        {"type": "text", "text": "封面", "page_idx": 8, "text_level": 1},
        {"type": "text", "text": "项目名称：XX 监理项目", "page_idx": 8},
        {"type": "text", "text": "第八章 投标文件格式", "page_idx": 9, "text_level": 1},
        {"type": "text", "text": "一、投标函", "page_idx": 9, "text_level": 2},
        {"type": "text", "text": "致：XX 招标人（投标函正文）", "page_idx": 9},
        {"type": "text", "text": "二、法定代表人身份证明", "page_idx": 10, "text_level": 2},
        {"type": "text", "text": "单位名称、姓名、职务等信息", "page_idx": 10},
        {
            "type": "table",
            "page_idx": 11,
            "table_body": (
                "<table><tr><th>姓名</th><th>职务</th></tr>"
                "<tr><td>张三</td><td>总监</td></tr></table>"
            ),
        },
    ]


class FakeMineruFormat(FakeMineru):
    async def __call__(self, input_file: Path, out_dir: Path, **kw: Any) -> RawResult:
        result = await super().__call__(input_file, out_dir, **kw)
        if input_file.stem in self.fail_stems:
            return result
        base = out_dir / input_file.stem
        stem = input_file.stem
        cl = base / f"{stem}_content_list.json"
        cl.write_text(json.dumps(_format_blocks(stem), ensure_ascii=False), encoding="utf-8")
        (base / f"{stem}.md").write_text("# 监理招标文件\n", encoding="utf-8")
        return RawResult(base, base / f"{stem}.md", cl, None, None, [])


@pytest.fixture
def project(db_client: TestClient) -> tuple[str, str]:
    ent = db_client.post("/api/v1/enterprises", json={"name": "测试企业", "agent": "张三"}).json()
    proj = db_client.post(
        f"/api/v1/enterprises/{ent['id']}/projects", json={"name": "测试项目"}
    ).json()
    return ent["id"], proj["id"]


def _register_and_parse(
    db_client: TestClient,
    eid: str,
    pid: str,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    file_name: str,
) -> list[dict[str, Any]]:
    src = tmp_path / file_name
    if file_name.endswith(".docx"):
        src.write_bytes(b"fake docx placeholder")  # fake MinerU 不读内容
    else:
        make_text_pdf(src, ["tender " * 20])
    db_client.post(
        f"/api/v1/enterprises/{eid}/projects/{pid}/parse/sources",
        json={"files": [{"name": file_name, "path": str(src)}]},
    )
    _patch_engine(monkeypatch, FakeMineruFormat(), {})
    events: list[dict[str, Any]] = []
    with db_client.stream(
        "POST",
        f"/api/v1/enterprises/{eid}/projects/{pid}/parse/start",
        json={"reparse": False},
    ) as resp:
        for line in resp.iter_lines():
            if line.startswith("data: "):
                events.append(json.loads(line[6:]))
    return events


def _assert_format_docx(eid: str, pid: str, stem: str) -> None:
    root = parse_service.paths.docx_dir(eid, pid)
    files = sorted(p.name for p in (root / stem).glob("*.docx"))
    # chapters 树 title 为去编号章名（编号在 number 字段），与 TR-12.4 示例一致
    assert files == ["00_封面.docx", "01_投标函.docx", "02_法定代表人身份证明.docx"]
    # 表格块已进入第二个子节 docx
    from docx import Document

    second = Document(root / stem / "02_法定代表人身份证明.docx")
    assert second.tables, "第二个子节应包含原表格"
    assert second.tables[0].rows[1].cells[0].text == "张三"


# ---------- TR-12.1：文字版 PDF ----------


def test_docx_pipeline_text_pdf(
    db_client: TestClient, project: tuple[str, str], tmp_path: Path, monkeypatch
) -> None:
    eid, pid = project
    events = _register_and_parse(db_client, eid, pid, tmp_path, monkeypatch, "监理招标文件.pdf")
    assert events[-1]["stage"] == "completed"

    _assert_format_docx(eid, pid, "监理招标文件")

    # manifest + 状态摘要
    manifest = json.loads(
        parse_service.paths.docx_manifest_path(eid, pid).read_text(encoding="utf-8")
    )
    assert manifest["completed"] is True
    assert manifest["total_files"] == 3
    assert manifest["sources"][0]["cover"] is True

    status = db_client.get(f"/api/v1/enterprises/{eid}/projects/{pid}/parse/status").json()
    assert status["parse_status"] == "SCORE_PARSED"
    assert status["docx"] == {
        "sources": 1,
        "files": 3,
        "cover": True,
        "completed": True,
        "errors": 0,
        "red_flags": 0,
        "missing_format_sources": [],
    }
    items = {i["key"]: i for i in status["checkpoint"]["items"]}
    assert items["docx:监理招标文件:00"]["state"] == "success"
    assert items["docx:监理招标文件:01"]["state"] == "success"
    assert items["docx:监理招标文件:02"]["state"] == "success"


# ---------- TR-12.2：扫描件 PDF（OCR content_list 同路径） ----------


def test_docx_pipeline_scanned_pdf(
    db_client: TestClient, project: tuple[str, str], tmp_path: Path, monkeypatch
) -> None:
    eid, pid = project
    events = _register_and_parse(db_client, eid, pid, tmp_path, monkeypatch, "扫描件招标文件.pdf")
    assert events[-1]["stage"] == "completed"
    _assert_format_docx(eid, pid, "扫描件招标文件")


# ---------- TR-12.3：docx 源文件 ----------


def test_docx_pipeline_docx_source(
    db_client: TestClient, project: tuple[str, str], tmp_path: Path, monkeypatch
) -> None:
    eid, pid = project
    events = _register_and_parse(db_client, eid, pid, tmp_path, monkeypatch, "监理招标文件.docx")
    assert events[-1]["stage"] == "completed"
    _assert_format_docx(eid, pid, "监理招标文件")


# ---------- TR-12.5：无格式章跳过不报错 ----------


def test_docx_without_format_chapter_skips(
    db_client: TestClient, project: tuple[str, str], tmp_path: Path, monkeypatch
) -> None:
    eid, pid = project
    src = tmp_path / "无格式章.pdf"
    make_text_pdf(src, ["tender " * 20])
    db_client.post(
        f"/api/v1/enterprises/{eid}/projects/{pid}/parse/sources",
        json={"files": [{"name": "无格式章.pdf", "path": str(src)}]},
    )
    _patch_engine(monkeypatch, FakeMineru(), {})
    events: list[dict[str, Any]] = []
    with db_client.stream(
        "POST",
        f"/api/v1/enterprises/{eid}/projects/{pid}/parse/start",
        json={"reparse": False},
    ) as resp:
        for line in resp.iter_lines():
            if line.startswith("data: "):
                events.append(json.loads(line[6:]))
    assert events[-1]["stage"] == "completed"
    assert any("未识别到「投标文件格式」章节" in e.get("message", "") for e in events)

    status = db_client.get(f"/api/v1/enterprises/{eid}/projects/{pid}/parse/status").json()
    # 决策 3A：空清单视为完成，不报错
    assert status["docx"]["files"] == 0
    assert status["docx"]["completed"] is True
    assert status["docx"]["missing_format_sources"] == ["无格式章"]


# ---------- TR-12.6：项目隔离 ----------


def test_docx_project_isolation(
    db_client: TestClient, project: tuple[str, str], tmp_path: Path, monkeypatch
) -> None:
    eid, pid = project
    _register_and_parse(db_client, eid, pid, tmp_path, monkeypatch, "监理招标文件.pdf")

    other_ent = db_client.post(
        "/api/v1/enterprises", json={"name": "其他企业", "agent": "李四"}
    ).json()
    other_proj = db_client.post(
        f"/api/v1/enterprises/{other_ent['id']}/projects", json={"name": "其他项目"}
    ).json()
    other_status = db_client.get(
        f"/api/v1/enterprises/{other_ent['id']}/projects/{other_proj['id']}/parse/status"
    ).json()
    assert other_status["docx"] is None
    # 越权访问本项目产物路径对应的 API（scope 校验 404）
    resp = db_client.get(f"/api/v1/enterprises/{other_ent['id']}/projects/{pid}/parse/status")
    assert resp.status_code == 404
    # 本企业目录下仅本项目有 docx 产物
    ent_root = parse_service.paths.project_dir(eid, pid)
    assert (ent_root / "parse" / "docx" / "监理招标文件").is_dir()


# ---------- TR-12.7：某章失败，重试只跑该章，成功章保留 ----------


def test_docx_chapter_failure_then_retry(
    db_client: TestClient, project: tuple[str, str], tmp_path: Path, monkeypatch
) -> None:
    eid, pid = project

    real_render = docx_build.render_section
    calls: list[int] = []

    def flaky_render(section, content_lists, project_dir, docx_root):  # noqa: ANN001
        calls.append(section.seq)
        if section.seq == 2 and not getattr(flaky_render, "healed", False):
            raise RuntimeError("simulated docx render failure")
        return real_render(section, content_lists, project_dir, docx_root)

    monkeypatch.setattr(docx_build, "render_section", flaky_render)
    events = _register_and_parse(db_client, eid, pid, tmp_path, monkeypatch, "监理招标文件.pdf")
    assert events[-1]["stage"] == "completed"

    status = db_client.get(f"/api/v1/enterprises/{eid}/projects/{pid}/parse/status").json()
    items = {i["key"]: i for i in status["checkpoint"]["items"]}
    assert items["docx:监理招标文件:02"]["state"] == "error"
    assert items["docx:监理招标文件:00"]["state"] == "success"
    assert items["docx:监理招标文件:01"]["state"] == "success"
    assert status["docx"]["completed"] is False
    assert status["docx"]["errors"] == 1

    root = parse_service.paths.docx_dir(eid, pid) / "监理招标文件"
    assert not (root / "02_法定代表人身份证明.docx").exists()
    cover_file = root / "00_封面.docx"
    cover_mtime = cover_file.stat().st_mtime_ns

    # 单项重试：先 reset 再续跑
    retry_resp = db_client.post(
        f"/api/v1/enterprises/{eid}/projects/{pid}/parse/retry",
        json={"item": "docx:监理招标文件:02"},
    )
    assert retry_resp.status_code == 200

    flaky_render.healed = True
    calls.clear()
    events2: list[dict[str, Any]] = []
    with db_client.stream(
        "POST",
        f"/api/v1/enterprises/{eid}/projects/{pid}/parse/start",
        json={"reparse": False},
    ) as resp:
        for line in resp.iter_lines():
            if line.startswith("data: "):
                events2.append(json.loads(line[6:]))
    assert events2[-1]["stage"] == "completed"

    # 续跑只渲染失败的 seq=2；成功章文件未重建
    assert calls == [2]
    assert cover_file.stat().st_mtime_ns == cover_mtime
    assert (root / "02_法定代表人身份证明.docx").exists()

    status2 = db_client.get(f"/api/v1/enterprises/{eid}/projects/{pid}/parse/status").json()
    assert status2["docx"]["completed"] is True
    assert status2["docx"]["errors"] == 0
    assert status2["docx"]["files"] == 3
